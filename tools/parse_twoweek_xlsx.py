"""Расписание «лист на группу, две недели в столбцах» (.xlsx) — в таблицу.

Так расписание ведёт, например, РГУ им. А.Н. Косыгина: на каждом листе —
одна группа; строки — день и пара (каждая пара двумя половинками по 40
минут); слева — нечётная неделя, справа — чётная:

    День | № | Время | Ауд. | Вид | Преподаватель | Нечётная | Чётная | Преподаватель | Вид | Ауд. | Время | №

Скрипт сводит это в простую таблицу CSV, которую принимает
tools/import_table.py, и печатает найденные звонки и срок семестра.

    python tools/parse_twoweek_xlsx.py файл.xlsx таблица.csv
"""

import csv
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

DAYS = {"ПН": "пн", "ВТ": "вт", "СР": "ср", "ЧТ": "чт", "ПТ": "пт", "СБ": "сб"}
TYPES = {"лек": "лекция", "пр": "практика", "лаб": "лабораторная", "сем": "семинар"}


def text(value):
    return re.sub(r"\s+", " ", str(value if value is not None else "")).strip()


def main():
    source, target = Path(sys.argv[1]), Path(sys.argv[2])
    book = load_workbook(source, data_only=True)
    rows, bells, period = [], {}, None
    for sheet in book.worksheets:
        group, course, header = "", 1, None
        for r in range(1, 40):
            line = " ".join(text(sheet.cell(r, c).value) for c in range(1, 15))
            found = re.search(r"ГРУППА\s+([\w-]+)", line)
            if found:
                group = found[1]
            found = re.search(r"(\d)\s*курс", line)
            if found:
                course = int(found[1])
            found = re.search(r"(\d{2})\.(\d{2})\.(\d{4})\s*-\s*(\d{2})\.(\d{2})\.(\d{4})", line)
            if found and not period:
                period = (f"{found[3]}-{found[2]}-{found[1]}", f"{found[6]}-{found[5]}-{found[4]}")
            if "День недели" in line:
                header = r
                break
        if not group or not header:
            print(f"лист «{sheet.title}»: не нашёл группу или шапку, пропускаю", file=sys.stderr)
            continue

        day = ""
        seen = set()
        count = 0
        for r in range(header + 1, sheet.max_row + 1):
            label = text(sheet.cell(r, 2).value).upper()
            if label in DAYS:
                day = DAYS[label]
            number = text(sheet.cell(r, 3).value)
            span = re.match(r"(\d{2}:\d{2})-(\d{2}:\d{2})", text(sheet.cell(r, 4).value))
            if not day or not number.isdigit() or not span:
                continue
            slot = int(number)
            # Пара записана двумя половинками: начало — у первой, конец — у второй.
            begin, finish = bells.get(slot, (span[1], span[2]))
            bells[slot] = (min(begin, span[1]), max(finish, span[2]))

            odd = (text(sheet.cell(r, 8).value), text(sheet.cell(r, 6).value), text(sheet.cell(r, 7).value), text(sheet.cell(r, 5).value))
            even = (text(sheet.cell(r, 9).value), text(sheet.cell(r, 11).value), text(sheet.cell(r, 10).value), text(sheet.cell(r, 12).value))
            variants = [("все", odd)] if odd == even else [("нечёт", odd), ("чёт", even)]
            for week, (subject, kind, teacher, room) in variants:
                if not subject or (day, slot, week, subject) in seen:
                    continue
                seen.add((day, slot, week, subject))
                count += 1
                rows.append([group, course, "бакалавриат", day, slot, week, subject, TYPES.get(kind.lower().rstrip("."), kind.lower()), teacher, room])
        print(f"лист «{sheet.title}»: группа {group}, {course} курс, пар {count}")

    with target.open("w", encoding="utf-8", newline="") as out:
        writer = csv.writer(out, delimiter=";")
        writer.writerow(["группа", "курс", "уровень", "день", "пара", "неделя", "предмет", "тип", "преподаватель", "аудитория"])
        writer.writerows(rows)
    print("звонки:", ",".join(f"{bells[n][0]}-{bells[n][1]}" for n in sorted(bells)))
    print("семестр:", *(period or ("?", "?")))
    print(f"строк: {len(rows)} → {target}")


if __name__ == "__main__":
    main()
