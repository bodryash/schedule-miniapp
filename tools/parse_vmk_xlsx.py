"""Расписание курса ВМК МГУ из Excel (.xlsx) — в папку факультета.

Таблица свободной вёрстки: лист на поток, шапка с группами повторяется перед
каждым днём, на пару отведено две строки. Обычно в верхней — предмет (и
аудитория), в нижней — преподаватель. Если же вся пара целиком («предмет
преподаватель аудитория») записана в одну из двух строк — это занятие через
неделю: верхняя строка — одна неделя, нижняя — другая. Какая из них чётная,
в файле не сказано, поэтому такие пары идут с пометкой «через неделю».

    python tools/parse_vmk_xlsx.py "3 курс.xlsx" --id msu-cmc --course 3 \\
        --from 2026-09-01 --to 2026-12-26 [--show]
"""

import argparse
import json
import re
import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "docs" / "data"

DAYS = {"понедельник": 1, "вторник": 2, "среда": 3, "четверг": 4, "пятница": 5, "суббота": 6}
TIME = re.compile(r"^\s*(\d{1,2})[.:](\d{2})\s*[-–—]+\s*(\d{1,2})[.:](\d{2})\s*$")
SURNAME = r"[А-ЯЁ][а-яё]+(?:-[А-ЯЁ][а-яё]+)?"
TITLE = r"(?:Чл\.-корр\.\s*РАН|профессор|прфессор|доцент|проф\.|доц\.)"
# «Иванов И.И.» (бывает «А..Н..») или «доцент Иванов Иван Иванович».
PERSON = re.compile(
    rf"(?:{TITLE}\s*)?{SURNAME}\s+[А-ЯЁ]\s*\.+\s*[А-ЯЁ]\s*\.+"
    rf"|{TITLE}\s+{SURNAME}\s+{SURNAME}\s+{SURNAME}"
    rf"|{SURNAME}\s+{SURNAME}\s+[А-ЯЁ][а-яё]+(?:вич|вна|ична|ьич)"
)
# Число — аудитория, если это не кусок даты или времени («17.09», «12:15»).
# Бывает приклеено к инициалам: «Фурсов А.С.607».
ROOM = re.compile(r"П-\d+а?|МЗ-\d+|ИВМ РАН|ЦФА|(?<![\d:])(?<!\d\.)\d{2,3}(?:-б)?\b(?![.:]\d)")
RENAME = {"ВП": "Военная подготовка", "Межфакультетские курсы": "Межфакультетские учебные курсы МГУ"}


def squash(text):
    return re.sub(r"\s+", " ", str(text or "").replace("\n", " ")).strip()


def clock(h, m):
    return f"{int(h):02d}:{m}"


def slot_of(start):
    minutes = int(start[:2]) * 60 + int(start[3:])
    for n, limit in enumerate((10 * 60, 12 * 60 + 10, 14 * 60 + 30, 16 * 60 + 15, 17 * 60 + 58), start=1):
        if minutes < limit:
            return n
    return 6


def split_rooms(text):
    rooms = [re.sub(r"\s+", " ", r) for r in ROOM.findall(text)]
    return re.sub(r"\s*,\s*(?=,|$)", "", ROOM.sub(" ", text)), ", ".join(dict.fromkeys(rooms))


def tidy(text):
    return re.sub(r"\s+([,.)])", r"\1", re.sub(r"\s+", " ", text)).strip(" ,;-–")


def parse_lesson(subject_text, teacher_text):
    """Предмет и строка преподавателя → поля пары и её особенности."""
    note, start, end = [], None, None
    text = subject_text
    # «с 21.09 …» — с какого числа; «с 13.30 …», «10.40 …», «10.30--13.25 …» — своё время.
    found = re.match(r"^с\s+(\d{1,2})[. ](\d{2})\.?\s+(?=\D)", text)
    if found and int(found[2]) <= 12 and int(found[1]) <= 31 and not (8 <= int(found[1]) <= 19 and found[2] in ("15", "30", "50", "40", "00")):
        note.append(f"с {int(found[1]):02d}.{found[2]}")
        text = text[found.end():]
    found = re.match(r"^(?:с\s+)?(\d{1,2})[.: ](\d{2})\.?\s*(?:[-–—]+\s*(\d{1,2})[.:](\d{2}))?\.?\s+(?=\D)", text)
    if found and 8 <= int(found[1]) <= 19:
        start = clock(found[1], found[2])
        end = clock(found[3], found[4]) if found[3] else None
        text = text[found.end():]
    for flag, meaning in (("с/с", "с/с"), ("д/в", "по выбору"), ("согл", "по согласованию")):
        if re.search(rf"(?<![а-яё]){re.escape(flag)}(?![а-яё])", text):
            note.append(meaning)
            text = re.sub(rf"(?<![а-яё]){re.escape(flag)}(?![а-яё])", " ", text)
    # «Экономика - 12» — номер группы по экономике, а не аудитория.
    found = re.search(r"(Экономика)\s*-\s*(\d+)", text)
    if found:
        note.append(f"группа по экономике {found[2]}")
        text = text[: found.start()] + found[1] + text[found.end():]
    text, room_a = split_rooms(text)
    teacher, room_b = split_rooms(teacher_text or "")
    subject = tidy(text)
    return {
        "subject": RENAME.get(subject, subject),
        "teacher": re.sub(r"\.{2,}", ".", tidy(teacher)).replace("прфессор", "профессор"),
        "room": ", ".join(r for r in (room_a, room_b) if r), "note": note, "start": start, "end": end,
    }


def whole(text):
    """Вся пара в одной строке? Тогда → (предмет, преподаватель)."""
    person = PERSON.search(text)
    if not person:
        return None
    head = tidy(ROOM.sub(" ", text[: person.start()]))
    return (text[: person.start()], text[person.start():]) if len(re.sub(r"[^А-Яа-яЁёA-Za-z]", "", head)) >= 3 else None


def read_sheet(sheet):
    merged = {}
    for rng in sheet.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = (rng.min_row, rng.min_col)
    value = lambda r, c: squash(sheet.cell(*merged.get((r, c), (r, c))).value)

    # Строки по порядку: шапка дня, начало пары или продолжение пары.
    kinds = {}
    for r in range(1, sheet.max_row + 1):
        label = squash(sheet.cell(r, 1).value).lower()
        if label in DAYS:
            kinds[r] = ("day", DAYS[label])
        elif TIME.match(label):
            kinds[r] = ("time", TIME.match(label).groups())
    lessons, groups, columns, day, odd = [], [], {}, 0, []
    for r in range(1, sheet.max_row + 1):
        kind = kinds.get(r)
        if not kind:
            continue
        if kind[0] == "day":
            day = kind[1]
            columns = {c: squash(sheet.cell(r, c).value) for c in range(2, sheet.max_column + 1)}
            columns = {c: name for c, name in columns.items() if re.fullmatch(r"\d{3}", name)}
            groups += [name for name in columns.values() if name not in groups]
            continue
        if not day or not columns:
            continue
        h1, m1, h2, m2 = kind[1]
        bell = (clock(h1, m1), clock(h2, m2))
        second = r + 1 if r + 1 <= sheet.max_row and r + 1 not in kinds else None
        last = max(columns)
        for c, group in columns.items():
            top = value(r, c)
            bottom = value(second, c) if second else ""
            # Хвост аудитории, уехавший в пустую соседнюю клетку за таблицей («64» + «7»).
            if c == last and second and re.fullmatch(r"\d", squash(sheet.cell(second, c + 1).value)):
                bottom += squash(sheet.cell(second, c + 1).value)
            if not top and not bottom:
                continue
            parts = []  # (предмет, преподаватель, через неделю?)
            top_whole, bottom_whole = whole(top) if top else None, whole(bottom) if bottom else None
            if top and bottom and not top_whole and not bottom_whole and not PERSON.search(top):
                # Нижняя строка повторяет слова предмета — это вторая пара, а не преподаватель.
                words = lambda text: {w.lower()[:6] for w in re.findall(r"[А-Яа-яЁё]{5,}", text)}
                if PERSON.search(bottom) or not (words(top) & words(bottom)):
                    parts.append((top, bottom, False))
                else:
                    # В обеих строках предмет без преподавателя — две пары через неделю.
                    parts += [(top, "", True), (bottom, "", True)]
            else:
                for text, split in ((top, top_whole), (bottom, bottom_whole)):
                    if not text:
                        continue
                    if split:
                        parts.append((split[0], split[1], True))
                    elif PERSON.search(text):
                        odd.append((sheet.title, r, group, text))  # одна фамилия без предмета
                    else:
                        parts.append((text, "", bool(top and bottom)))
            for subject_text, teacher_text, half in parts:
                item = parse_lesson(subject_text, teacher_text)
                if not item["subject"]:
                    odd.append((sheet.title, r, group, subject_text))
                    continue
                start, end = item.pop("start") or bell[0], item.pop("end")
                note = item.pop("note") + (["через неделю"] if half else [])
                lessons.append({
                    "group": group, "day": day, "slot": slot_of(bell[0]), "subgroup": None, "elective": None, "week": "all",
                    "note": "; ".join(note), "link": "", "start": start, "end": end or bell[1], "type": "", **item,
                })
    return groups, lessons, odd


def main():
    parser = argparse.ArgumentParser(description="Расписание курса ВМК МГУ из .xlsx")
    parser.add_argument("table", type=Path)
    parser.add_argument("--id", required=True)
    parser.add_argument("--university", default="МГУ")
    parser.add_argument("--faculty", default="Факультет вычислительной математики и кибернетики")
    parser.add_argument("--short", default="ВМК")
    parser.add_argument("--course", type=int, required=True)
    parser.add_argument("--level", default="бакалавриат")
    parser.add_argument("--from", dest="start", required=True)
    parser.add_argument("--to", dest="end", required=True)
    parser.add_argument("--show", action="store_true", help="только напечатать разобранное, ничего не записывать")
    args = parser.parse_args()
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{1,30}", args.id) or args.id == "msu-fgp":
        sys.exit("--id: латиница, цифры и дефис; msu-fgp занят")

    names, lessons, odd = [], [], []
    for sheet in openpyxl.load_workbook(args.table).worksheets:
        found = read_sheet(sheet)
        names += [g for g in found[0] if g not in names]
        lessons += found[1]
        odd += found[2]

    # Одна и та же пара у соседних групп: аудиторию берём самую полную («64» и «647»).
    best = {}
    for lesson in lessons:
        key = (lesson["day"], lesson["start"], lesson["subject"], lesson["teacher"])
        if len(lesson["room"]) > len(best.setdefault(key, "")):
            best[key] = lesson["room"]
    for lesson in lessons:
        full = best[(lesson["day"], lesson["start"], lesson["subject"], lesson["teacher"])]
        if full.startswith(lesson["room"]):
            lesson["room"] = full
    # Пара на два звонка со своим временем («10.30--13.25») записана дважды — оставляем одну.
    seen, unique = set(), []
    for lesson in sorted(lessons, key=lambda l: (l["group"], l["day"], l["slot"])):
        key = (lesson["group"], lesson["day"], lesson["start"], lesson["end"], lesson["subject"], lesson["teacher"], lesson["note"])
        if key not in seen:
            seen.add(key)
            unique.append(lesson)
    lessons = unique

    for title, row, group, text in odd:
        print(f"не разобрал: лист «{title}», строка {row}, группа {group}: {text}", file=sys.stderr)
    if args.show:
        day_names = "_ Пн Вт Ср Чт Пт Сб".split()
        for l in sorted(lessons, key=lambda l: (l["group"], l["day"], l["start"])):
            print(f"{l['group']} {day_names[l['day']]} {l['slot']} {l['start']}-{l['end']} | {l['subject']} | {l['teacher']} | {l['room']} | {l['note']}")
        print(f"групп {len(names)}, пар {len(lessons)}, не разобрано {len(odd)}")
        return
    if odd or not lessons:
        sys.exit("Разобрано не всё — ничего не записываю.")

    groups = [{"id": name, "title": name, "course": args.course, "level": args.level} for name in names]
    start, end = date.fromisoformat(args.start), date.fromisoformat(args.end)
    weeks, cursor, parity = [], start, "odd"
    while cursor <= end:
        sunday = cursor + timedelta(days=6 - cursor.weekday())
        weeks.append({"from": cursor.isoformat(), "to": min(sunday, end).isoformat(), "parity": parity})
        cursor, parity = sunday + timedelta(days=1), ("even" if parity == "odd" else "odd")
    bells = [{"n": n, "start": s, "end": e} for n, (s, e) in enumerate(
        [("08:45", "10:20"), ("10:30", "12:05"), ("12:50", "14:25"), ("14:35", "16:10"), ("16:20", "17:55"), ("18:00", "19:30")], start=1)]

    folder = DATA / args.id
    folder.mkdir(parents=True, exist_ok=True)
    target = folder / "schedule.json"
    schedule = {
        "meta": {"updated": date.today().isoformat(), "semester": "", "faculty": f"{args.faculty} {args.university}",
                 "period": {"from": start.strftime("%d.%m.%Y"), "to": end.strftime("%d.%m.%Y")}},
        "weeks": weeks, "bells": bells, "groups": groups, "lessons": lessons,
    }
    # Другие курсы факультета, подключённые раньше, остаются на месте.
    if target.exists():
        old = json.loads(target.read_text(encoding="utf-8"))
        mine = set(names)
        schedule["groups"] = [g for g in old.get("groups", []) if g["id"] not in mine] + groups
        schedule["lessons"] = [l for l in old.get("lessons", []) if l["group"] not in mine] + lessons
    schedule["lessons"].sort(key=lambda l: (l["group"], l["day"], l["start"], l["slot"]))
    dump = lambda value: json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    target.write_text(dump(schedule), encoding="utf-8")
    (folder / "groups.json").write_text(dump(schedule["groups"]), encoding="utf-8")

    registry = DATA / "tenants.json"
    tenants = json.loads(registry.read_text(encoding="utf-8"))
    keep = next((t for t in tenants if t["id"] == args.id), {})
    entry = {**keep, "id": args.id, "university": args.university, "faculty": args.faculty, "path": args.id,
             "features": keep.get("features", []), "short": keep.get("short", args.short)}
    registry.write_text(json.dumps([t for t in tenants if t["id"] != args.id] + [entry], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Итого: групп {len(groups)}, пар {len(lessons)}, записано в docs/data/{args.id}/")


if __name__ == "__main__":
    main()
