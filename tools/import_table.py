"""Расписание факультета из таблицы — в папку приложения.

У каждого вуза свой формат расписания, и писать разборщик под каждый PDF
долго. Поэтому вход здесь — простая таблица (CSV или Excel), в которую
расписание можно свести руками или скриптом. Одна строка — одна пара:

    группа, курс, уровень, день, пара, неделя, предмет, тип, преподаватель,
    аудитория, подгруппа, начало, конец, примечание

Обязательны только группа, день, пара и предмет. День — «пн»…«сб» или 1…6.
Неделя — «все», «чёт», «нечёт» (пусто = все). Уровень — «бакалавриат»
(по умолчанию), «магистратура», «специалитет».

Запуск:

    python tools/import_table.py таблица.csv --id spbu-law \\
        --university "СПбГУ" --faculty "Юридический факультет" \\
        --from 2026-09-01 --to 2026-12-27 --first-week odd

Скрипт создаёт docs/data/<id>/schedule.json и groups.json и вписывает
факультет в docs/data/tenants.json. Повторный запуск обновляет расписание.
"""

import argparse
import csv
import json
import re
import sys
from datetime import date, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "docs" / "data"

# Звонки по умолчанию — как на ФГП МГУ; свои задаются через --bells.
DEFAULT_BELLS = "09:00-10:30,10:45-12:15,13:00-14:30,14:45-16:15,16:30-18:00,18:10-19:40,19:50-21:20"

COLUMNS = {
    "group": ["группа", "group"],
    "course": ["курс", "course"],
    "level": ["уровень", "ступень", "level"],
    "day": ["день", "день недели", "day"],
    "slot": ["пара", "номер пары", "slot", "№"],
    "week": ["неделя", "чётность", "четность", "week"],
    "subject": ["предмет", "дисциплина", "subject"],
    "type": ["тип", "вид", "вид занятия", "type"],
    "teacher": ["преподаватель", "teacher"],
    "room": ["аудитория", "ауд", "ауд.", "room"],
    "subgroup": ["подгруппа", "subgroup"],
    "start": ["начало", "start"],
    "end": ["конец", "end"],
    "note": ["примечание", "note"],
}
DAYS = {"пн": 1, "вт": 2, "ср": 3, "чт": 4, "пт": 5, "сб": 6,
        "понедельник": 1, "вторник": 2, "среда": 3, "четверг": 4, "пятница": 5, "суббота": 6,
        "mon": 1, "tue": 2, "wed": 3, "thu": 4, "fri": 5, "sat": 6}
WEEKS = {"": "all", "все": "all", "all": "all", "каждая": "all",
         "чёт": "even", "чет": "even", "чётная": "even", "четная": "even", "even": "even",
         "нечёт": "odd", "нечет": "odd", "нечётная": "odd", "нечетная": "odd", "odd": "odd"}


def read_rows(path):
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        try:
            from openpyxl import load_workbook
        except ImportError:
            sys.exit("Для Excel нужен пакет openpyxl: pip install openpyxl. Или сохраните таблицу как CSV.")
        sheet = load_workbook(path, read_only=True, data_only=True).active
        rows = [["" if cell is None else str(cell).strip() for cell in row] for row in sheet.iter_rows(values_only=True)]
    else:
        text = path.read_text(encoding="utf-8-sig")
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t")
        rows = [[cell.strip() for cell in row] for row in csv.reader(text.splitlines(), dialect)]
    rows = [row for row in rows if any(row)]
    if not rows:
        sys.exit("Таблица пустая.")
    head = [cell.lower() for cell in rows[0]]
    index = {}
    for key, names in COLUMNS.items():
        for name in names:
            if name in head:
                index[key] = head.index(name)
                break
    missing = [COLUMNS[key][0] for key in ("group", "day", "slot", "subject") if key not in index]
    if missing:
        sys.exit(f"В первой строке таблицы нет столбцов: {', '.join(missing)}. Есть: {', '.join(rows[0])}")
    return [{key: (row[i] if i < len(row) else "") for key, i in index.items()} for row in rows[1:]]


def build_weeks(start, end, first):
    """Учебные недели с чередованием чётности, от понедельника до воскресенья."""
    weeks = []
    cursor = start
    parity = first
    while cursor <= end:
        sunday = cursor + timedelta(days=6 - cursor.weekday())
        weeks.append({"from": cursor.isoformat(), "to": min(sunday, end).isoformat(), "parity": parity})
        cursor = sunday + timedelta(days=1)
        parity = "even" if parity == "odd" else "odd"
    return weeks


def main():
    parser = argparse.ArgumentParser(description="Расписание факультета из таблицы")
    parser.add_argument("table", type=Path)
    parser.add_argument("--id", required=True, help="латиницей: spbu-law")
    parser.add_argument("--university", required=True)
    parser.add_argument("--faculty", required=True)
    parser.add_argument("--from", dest="start", required=True, help="первый день семестра, 2026-09-01")
    parser.add_argument("--to", dest="end", required=True, help="последний день занятий")
    parser.add_argument("--first-week", choices=["odd", "even"], default="odd", help="чётность первой недели")
    parser.add_argument("--bells", default=DEFAULT_BELLS, help="звонки: 09:00-10:30,10:45-12:15,…")
    args = parser.parse_args()

    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{1,30}", args.id):
        sys.exit("--id: только латиница, цифры и дефис, например spbu-law")
    if args.id == "msu-fgp":
        sys.exit("msu-fgp — это ФГП, его расписание собирается из PDF отдельно.")

    bells = []
    for n, pair in enumerate(args.bells.split(","), 1):
        begin, finish = pair.strip().split("-")
        bells.append({"n": n, "start": begin, "end": finish})

    lessons, groups, problems = [], {}, []
    for line, row in enumerate(read_rows(args.table), 2):
        group = row.get("group", "")
        subject = row.get("subject", "")
        if not group and not subject:
            continue
        day = DAYS.get(row.get("day", "").lower()) or (int(row["day"]) if row.get("day", "").isdigit() else 0)
        slot = int(row["slot"]) if row.get("slot", "").isdigit() else 0
        week = WEEKS.get(row.get("week", "").lower())
        if not group or not subject or not 1 <= day <= 6 or not 1 <= slot <= len(bells) or week is None:
            problems.append(f"строка {line}: {' | '.join(v for v in row.values() if v)[:90]}")
            continue
        course = int(row["course"]) if row.get("course", "").isdigit() else 1
        level = row.get("level", "").lower() or "бакалавриат"
        groups.setdefault(group, {"id": group, "title": group, "course": course, "level": level})
        subgroup = int(row["subgroup"]) if row.get("subgroup", "").isdigit() else None
        lessons.append({
            "group": group, "day": day, "slot": slot, "subgroup": subgroup, "elective": None, "week": week,
            "note": row.get("note", ""), "link": "", "start": row.get("start", ""), "end": row.get("end", ""),
            "type": row.get("type", "").lower(), "subject": subject, "room": row.get("room", ""),
            "teacher": row.get("teacher", ""),
        })

    if problems:
        print(f"Не разобрал строк: {len(problems)}", file=sys.stderr)
        for problem in problems[:15]:
            print("  " + problem, file=sys.stderr)
        if len(problems) > len(lessons) / 10:
            sys.exit("Ошибок слишком много — расписание не записано. Проверьте столбцы день, пара и неделя.")
    if not lessons:
        sys.exit("Ни одной пары не разобрано.")

    start, end = date.fromisoformat(args.start), date.fromisoformat(args.end)
    group_list = sorted(groups.values(), key=lambda g: (g["level"], g["course"], g["id"]))
    schedule = {
        "meta": {
            "updated": date.today().isoformat(),
            "semester": "",
            "faculty": f"{args.faculty} {args.university}",
            "period": {"from": start.strftime("%d.%m.%Y"), "to": end.strftime("%d.%m.%Y")},
        },
        "weeks": build_weeks(start, end, args.first_week),
        "bells": bells,
        "groups": group_list,
        "lessons": sorted(lessons, key=lambda l: (l["group"], l["day"], l["slot"])),
    }

    folder = DATA / args.id
    folder.mkdir(parents=True, exist_ok=True)
    dump = lambda value: json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    (folder / "schedule.json").write_text(dump(schedule), encoding="utf-8")
    (folder / "groups.json").write_text(dump(group_list), encoding="utf-8")

    registry = DATA / "tenants.json"
    tenants = json.loads(registry.read_text(encoding="utf-8"))
    keep = next((t for t in tenants if t["id"] == args.id), {})
    entry = {**keep, "id": args.id, "university": args.university, "faculty": args.faculty, "path": args.id, "features": keep.get("features", [])}
    tenants = [t for t in tenants if t["id"] != args.id] + [entry]
    registry.write_text(json.dumps(tenants, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    print(f"{args.university}, {args.faculty}: групп {len(group_list)}, пар {len(lessons)}, недель {len(schedule['weeks'])}")
    print(f"Записано в docs/data/{args.id}/")


if __name__ == "__main__":
    main()
