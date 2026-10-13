"""Расписание географического факультета МГУ из Excel (.xlsx) — в папку факультета.

Одна книга на весь факультет: лист на курс, столбцы — группы (на первом
курсе) или кафедры (дальше), шапка со столбцами повторяется перед каждым
днём. Пара бывает разделена на две строки, «верх» и «нижн»: клетка на обе
строки идёт каждую неделю, клетка в одной из них — через неделю. Верхняя
неделя — нечётная: в таблице на неё приходится 12 октября.

Текст клетки свободный («Математика 2012 Пр. Шарапова»). Достаём из него то,
что узнаётся наверняка — вид занятия, преподавателя, аудиторию, — остальное
оставляем как написано. Ссылки на онлайн-занятия не публикуем: таблица
факультета закрыта для посторонних.

    python tools/parse_geo_xlsx.py расписание.xlsx --id msu-geo \\
        --from 2026-09-01 --to 2026-12-26 [--show]
"""

import argparse
import json
import re
import sys
from datetime import date, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "docs" / "data"

DAY_ORDER = ["понедельник", "вторник", "среда", "четверг", "пятница", "суббота"]
TIME = re.compile(r"(\d{1,2})\s*[-.:]\s*(\d{2})\s*[-–—]+\s*(\d{1,2})\s*[-.:]\s*(\d{2})")
URL = re.compile(r"https?://\S+")
KINDS = [(re.compile(r"\(\s*[лЛ]\s*\)"), "лекция"), (re.compile(r"\(\s*[сС]\s*\)"), "семинар"),
         (re.compile(r"\(\s*пр\s*\)", re.I), "практика"), (re.compile(r"\(\s*практикум\s*\)", re.I), "практикум")]
TITLE = r"(?:проф|прф|доц|ст\.?\s*пр(?:еп)?|преп|пр|асс|акад|внс|снс|мнс|нс)"
NAME = r"[А-ЯЁC][А-Яа-яёЁ-]{2,}"
TEACHER = re.compile(rf"(?<![А-Яа-яЁё])((?i:{TITLE}))(?:\.\s*|\s+)({NAME})((?:\s*,\s*(?:(?i:{TITLE})(?:\.\s*|\s+))?{NAME})*)")
UP = r"(?:/\s*)?(?:по\s+)?(?:верх\w*\.{0,2}\s*нед\w*\.?|в\.\s*нед\.?)"
DOWN = r"(?:/\s*)?(?:по\s+)?нижн\w*\.{0,2}\s*нед\w*\.?"
MONTHS = "января|февраля|марта|апреля|мая|июня|сентября|октября|ноября|декабря"
SKIP = re.compile(r"^день самостоятельной работы$", re.I)


def squash(text):
    return re.sub(r"\s+", " ", str(text or "")).strip()


def unspace(text):
    """«Г И Д Р О Л О Г И Я  (Л)» → «ГИДРОЛОГИЯ (Л)»: буквы вразрядку, слова через два пробела."""
    parts = re.split(r"\s{2,}", str(text or "").replace("\n", "  "))
    joined = [p.replace(" ", "") if re.fullmatch(r"(?:[А-ЯЁ] )+[А-ЯЁ]", p.strip()) else p.strip() for p in parts]
    return " ".join(p for p in joined if p)


def sentence(text):
    """ЗАГЛАВНЫЕ → обычные; короткие сокращения вроде «МФК» не трогаем."""
    lowered = re.sub(r"[А-ЯЁ]{4,}", lambda m: m[0].lower(), text)
    if lowered != text:
        keep = {"МФК", "ГУМ", "РАН", "ФГМ", "ГИС", "ДЗЗ", "РФ", "ПП", "НИР", "ОС", "ИС", "ЭВМ"}
        lowered = re.sub(r"(?<![А-Яа-яЁё])[А-ЯЁ]{1,3}(?![А-Яа-яЁё.])", lambda m: m[0] if m[0] in keep else m[0].lower(), lowered)
        text = lowered[:1].upper() + lowered[1:]
        text = re.sub(r"\bроссии\b", "России", text)
    return text


def person(match):
    """«ПРОФ. ФРОЛОВА» → «проф. Фролова»; несколько фамилий через запятую — так же."""
    def name(word):
        word = word.replace("C", "С")  # латинская C в «CЕМЕНОВ»
        return word[:1].upper() + word[1:].lower() if word.isupper() else word
    title = re.sub(r"\s+", " ", match[1].lower()).replace("прф", "проф")
    dot = lambda t: t if t in ("нс", "внс", "снс", "мнс") else f"{t}."
    first = f"{dot(title)} {name(match[2])}"
    rest = [part.strip() for part in match[3].split(",") if part.strip()]
    more = []
    for part in rest:
        found = re.match(rf"((?i:{TITLE}))(?:\.\s*|\s+)({NAME})$", part)
        more.append(f"{dot(found[1].lower())} {name(found[2])}" if found else name(part))
    return ", ".join([first] + more)


def parse_text(raw):
    """Текст клетки → поля пары."""
    text = unspace(raw)
    note = []
    online = bool(URL.search(text))
    text = URL.sub(" ", text)
    for found in re.finditer(rf"\b(\d{{1,2}})\s+({MONTHS})\b", text):
        note.append(f"{found[1]} {found[2]} — онлайн, ссылка в расписании факультета" if online else f"{found[1]} {found[2]}")
    text = re.sub(rf"\b\d{{1,2}}\s+(?:{MONTHS})\b", " ", text)
    kind = ""
    for pattern, name in KINDS:
        if pattern.search(text):
            kind = kind or name
            text = pattern.sub(" ", text)
    if re.search(r"\b3\s*часа\b", text):
        note.append("3 часа")
        text = re.sub(r"\b3\s*часа\b", " ", text)
    text = squash(text)
    begin = re.match(r"^[сС]\s+(\d{1,2})(?:\s*[-.:]\s*(\d{2}))?\s*(?:часов|ч\.*)?\s+(?=[А-ЯЁа-яё])", text)
    start = None
    if begin and 8 <= int(begin[1]) <= 20:
        start = f"{int(begin[1]):02d}:{begin[2] or '00'}"
        text = text[begin.end():]

    # Подгруппы с разными предметами и аудиториями — оставляем как написано.
    if re.search(r"п/[гп]", text):
        return {"subject": sentence(text), "type": kind, "teacher": "", "room": "", "note": "; ".join(note), "own": start}
    # «Иностранный язык (а201 1917) (а202 2005) …» — название и список групп с аудиториями.
    lang = re.match(r"^(Иностранный язык)\s*(.*)$", text, re.I)
    if lang and lang[2] and not TEACHER.search(lang[2]) and re.search(r"\(|[а-я]\d", lang[2]):
        note.insert(0, lang[2])
        return {"subject": "Иностранный язык", "type": kind, "teacher": "", "room": "", "note": "; ".join(note), "own": start}

    teachers = []
    def take(match):
        teachers.append(person(match))
        return " "
    text = TEACHER.sub(take, text)
    rooms = []
    def room(match):
        rooms.append(squash(match[0]))
        return " "
    text = re.sub(r"[12]\s*ГУМ\s*(?:ауд\.?\s*)?\d{3,4}|\d{3,4}\s*[12]\s*ГУМ", room, text)
    text = re.sub(r"(?<![\d.:-])\b(?:\d{4}[а-я]?|02|\d{3})\b(?![.:-]\d)", room, text)
    subject = sentence(squash(text).strip(" ,.;-"))
    subject = subject[:1].upper() + subject[1:]
    return {"subject": subject, "type": kind, "teacher": ", ".join(teachers), "room": ", ".join(rooms), "note": "; ".join(note), "own": start}


def split_weeks(raw, week):
    """«верхн.нед. X  нижн.нед. Y» в одной клетке → две пары по своим неделям."""
    text = unspace(raw)
    both = re.match(rf"^\s*{UP}\s*(.+?)\s*{DOWN}\s*(.+)$", text, re.I | re.S)
    if both:
        return [(both[1], "odd"), (both[2], "even")]
    upper = re.match(rf"^\s*{UP}\s*(.+)$", text, re.I | re.S)
    if upper and not re.search(DOWN, upper[1], re.I):
        return [(upper[1], "odd")]
    lower = re.match(rf"^\s*{DOWN}\s*(.+)$", text, re.I | re.S)
    if lower and not re.search(UP, lower[1], re.I):
        return [(lower[1], "even")]
    ups, downs = len(re.findall(UP, text, re.I)), len(re.findall(DOWN, text, re.I))
    if ups == 1 and not downs:
        return [(re.sub(UP, " ", text, flags=re.I), "odd")]
    if downs == 1 and not ups:
        return [(re.sub(DOWN, " ", text, flags=re.I), "even")]
    return [(text, week)]


def read_sheet(sheet, course, level, suffix):
    merged = {}
    for rng in sheet.merged_cells.ranges:
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                merged[(r, c)] = (rng.min_row, rng.min_col, rng.max_row, rng.max_col)
    cell = lambda r, c: sheet.cell(r, c).value

    # Шапка: строка, где в столбце C «Пара», либо первая, где в D стоит номер группы.
    header = None
    for r in range(1, 6):
        c3, d = squash(cell(r, 3)).lower(), squash(cell(r, 4))
        if c3 == "пара" or re.fullmatch(r"\d{3}(?:\.0)?", d):
            header = r
            break
    if not header:
        return [], [], [f"лист «{sheet.title}»: шапка не найдена"]
    key = squash(cell(header, 4))
    header_rows = [r for r in range(header, sheet.max_row + 1) if squash(cell(r, 4)) == key and not TIME.search(squash(cell(r, 2)))]
    names = {}
    for r in header_rows:
        for c in range(4, sheet.max_column + 1):
            value = squash(cell(r, c))
            if value and value.lower() != "пара" and len(value) < 40 and c not in names:
                names[c] = re.sub(r"\.0$", "", value)
    taken, columns = set(), {}
    for c in sorted(names):
        name = names[c] if not suffix or re.fullmatch(r"\d{3}", names[c]) else f"{names[c]}{suffix}"
        if re.fullmatch(r"Туризм", name):
            name = f"Туризм{suffix or f'-{course}'}"
        while name in taken:
            name += " (2)"
        taken.add(name)
        columns[c] = name

    lessons, odd, done = [], [], set()
    day, time, slot = -1, None, 0
    for r in range(header, sheet.max_row + 1):
        if r in header_rows:
            if r == header or r - 1 not in header_rows:
                day += 1
            continue
        found = TIME.search(squash(cell(r, 2)))
        if found:
            time = (f"{int(found[1]):02d}:{found[2]}", f"{int(found[3]):02d}:{found[4]}")
        label = squash(cell(r, 3)).lower()
        number = re.match(r"(\d)", label)
        if number:
            slot = int(number[1])
        if not label and not found:
            continue
        if day < 0 or day > 5 or not time or not slot:
            continue
        half = "odd" if "верх" in label else "even" if "нижн" in label else None
        for c, group in columns.items():
            r0, c0, r1, c1 = merged.get((r, c), (r, c, r, c))
            raw = cell(r0, c0)
            if raw is None or not str(raw).strip() or (r0, c0, group) in done:
                continue
            done.add((r0, c0, group))
            # Клетка на обе строки пары — каждую неделю; в одной из них — через неделю.
            week = "all" if (r1 > r0 or not half) else half
            for text, when in split_weeks(str(raw), week):
                if SKIP.match(squash(unspace(text))):
                    continue
                item = parse_text(text)
                if not item["subject"]:
                    odd.append(f"лист «{sheet.title}», строка {r0}, {group}: {squash(raw)[:80]}")
                    continue
                own = item.pop("own")
                begin, finish = own or time[0], time[1]
                if own and own != time[0]:
                    total = int(begin[:2]) * 60 + int(begin[3:]) + (135 if "3 часа" in item["note"] else 95)
                    finish = f"{total // 60:02d}:{total % 60:02d}"
                lessons.append({
                    "group": group, "day": day + 1, "slot": slot, "subgroup": None, "elective": None, "week": when,
                    "link": "", "start": begin, "end": finish, **item,
                })
    groups = [{"id": name, "title": name, "course": course, "level": level} for name in columns.values()]
    if day != 5:
        odd.append(f"лист «{sheet.title}»: дней найдено {day + 1}, а не 6")
    return groups, lessons, odd


def main():
    parser = argparse.ArgumentParser(description="Расписание геофака МГУ из .xlsx")
    parser.add_argument("table", type=Path)
    parser.add_argument("--id", required=True)
    parser.add_argument("--university", default="МГУ")
    parser.add_argument("--faculty", default="Географический факультет")
    parser.add_argument("--short", default="Геофак")
    parser.add_argument("--from", dest="start", required=True)
    parser.add_argument("--to", dest="end", required=True)
    parser.add_argument("--show", action="store_true")
    args = parser.parse_args()
    if not re.fullmatch(r"[a-z0-9][a-z0-9-]{1,30}", args.id) or args.id == "msu-fgp":
        sys.exit("--id: латиница, цифры и дефис; msu-fgp занят")

    groups, lessons, odd = [], [], []
    for sheet in openpyxl.load_workbook(args.table).worksheets:
        title = sheet.title.lower()
        course = re.search(r"(\d)\s*курс", title)
        master = re.search(r"магистры\s*(\d)", title)
        if course:
            found = read_sheet(sheet, int(course[1]), "бакалавриат", "" if course[1] == "1" else f"-{course[1]}")
        elif master:
            found = read_sheet(sheet, int(master[1]), "магистратура", f"-м{master[1]}")
        else:
            print(f"лист «{sheet.title}» пропущен", file=sys.stderr)
            continue
        groups += found[0]
        lessons += found[1]
        odd += found[2]
    # Группы без единой пары (пустой столбец) не показываем.
    busy = {l["group"] for l in lessons}
    groups = [g for g in groups if g["id"] in busy]
    for line in odd:
        print("не разобрал:", line, file=sys.stderr)
    lessons.sort(key=lambda l: (l["group"], l["day"], l["start"], l["slot"]))
    if args.show:
        names = "_ Пн Вт Ср Чт Пт Сб".split()
        for l in lessons:
            print(f"{l['group']} | {names[l['day']]} {l['slot']} {l['start']}-{l['end']} {l['week']} | {l['subject']} | {l['type']} | {l['teacher']} | {l['room']} | {l['note']}")
        print(f"групп {len(groups)}, пар {len(lessons)}, не разобрано {len(odd)}")
        return
    if odd or not lessons:
        sys.exit("Разобрано не всё — ничего не записываю.")

    start, end = date.fromisoformat(args.start), date.fromisoformat(args.end)
    weeks, cursor, parity = [], start, "odd"
    while cursor <= end:
        sunday = cursor + timedelta(days=6 - cursor.weekday())
        weeks.append({"from": cursor.isoformat(), "to": min(sunday, end).isoformat(), "parity": parity})
        cursor, parity = sunday + timedelta(days=1), ("even" if parity == "odd" else "odd")
    bells = [{"n": n, "start": s, "end": e} for n, (s, e) in enumerate(
        [("09:00", "10:35"), ("10:50", "12:25"), ("13:10", "14:45"), ("14:55", "16:30"), ("16:40", "18:15"), ("18:25", "20:00")], start=1)]
    schedule = {
        "meta": {"updated": date.today().isoformat(), "semester": "", "faculty": f"{args.faculty} {args.university}",
                 "period": {"from": start.strftime("%d.%m.%Y"), "to": end.strftime("%d.%m.%Y")}},
        "weeks": weeks, "bells": bells, "groups": groups, "lessons": lessons,
    }
    folder = DATA / args.id
    folder.mkdir(parents=True, exist_ok=True)
    dump = lambda value: json.dumps(value, ensure_ascii=False, separators=(",", ":"))
    (folder / "schedule.json").write_text(dump(schedule), encoding="utf-8")
    (folder / "groups.json").write_text(dump(groups), encoding="utf-8")
    registry = DATA / "tenants.json"
    tenants = json.loads(registry.read_text(encoding="utf-8"))
    keep = next((t for t in tenants if t["id"] == args.id), {})
    entry = {**keep, "id": args.id, "university": args.university, "faculty": args.faculty, "path": args.id,
             "features": keep.get("features", []), "short": keep.get("short", args.short)}
    registry.write_text(json.dumps([t for t in tenants if t["id"] != args.id] + [entry], ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Итого: групп {len(groups)}, пар {len(lessons)}, записано в docs/data/{args.id}/")


if __name__ == "__main__":
    main()
